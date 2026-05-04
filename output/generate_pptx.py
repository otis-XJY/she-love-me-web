#!/usr/bin/env python3
"""
Generate PPTX for: She Loves Me, She Loves Me Not
Interactive Digital Poetry Installation — Psychology Project Product Report
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Constants ──────────────────────────────────────────────
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Colors
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x2D, 0x2D, 0x2D)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MID_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
ACCENT = RGBColor(0xC0, 0x6C, 0x84)  # Rose — matches the poem/flower theme
ACCENT2 = RGBColor(0x6B, 0x8E, 0xA3)  # Muted blue
BG_WHITE = RGBColor(0xFA, 0xFA, 0xFA)

FONT_TITLE = "Microsoft YaHei"
FONT_BODY = "Microsoft YaHei"
FONT_EN = "Calibri"

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


# ── Helpers ────────────────────────────────────────────────
def add_blank_slide():
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    # white background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return slide


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=BLACK, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name=FONT_BODY, line_spacing=1.4):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.line_spacing = Pt(int(font_size * line_spacing))
    return txBox


def add_multiline_box(slide, left, top, width, height, lines, font_size=16,
                      color=BLACK, bold=False, bullet=False,
                      alignment=PP_ALIGN.LEFT, line_spacing=1.5):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        if bullet:
            p.text = "•  " + line
        else:
            p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = FONT_BODY
        p.alignment = alignment
        p.line_spacing = Pt(int(font_size * line_spacing))
    return txBox


def add_accent_bar(slide, left, top, width=Inches(0.08), height=Inches(0.6)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    return shape


def add_slide_number(slide, num, total):
    add_text_box(slide, Inches(12.4), Inches(7.0), Inches(0.8), Inches(0.4),
                 f"{num}/{total}", font_size=10, color=MID_GRAY,
                 alignment=PP_ALIGN.RIGHT)


def add_source_label(slide, text, left=Inches(0.6), top=Inches(6.9)):
    add_text_box(slide, left, top, Inches(6), Inches(0.4),
                 text, font_size=9, color=MID_GRAY)


def add_section_title(slide, title, subtitle=None):
    add_accent_bar(slide, Inches(0.6), Inches(1.5), Inches(0.08), Inches(0.7))
    add_text_box(slide, Inches(0.9), Inches(1.5), Inches(10), Inches(0.8),
                 title, font_size=32, color=DARK_GRAY, bold=True)
    if subtitle:
        add_text_box(slide, Inches(0.9), Inches(2.3), Inches(10), Inches(0.5),
                     subtitle, font_size=16, color=MID_GRAY)


TOTAL_SLIDES = 13


# ══════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
# Decorative accent strip at top
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.12))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()

add_text_box(s, Inches(0.8), Inches(1.8), Inches(11), Inches(1.2),
             "She Loves Me, She Loves Me Not",
             font_size=40, color=DARK_GRAY, bold=True, font_name=FONT_EN)

add_text_box(s, Inches(0.8), Inches(3.0), Inches(11), Inches(0.8),
             "人类与AI的情感纠缠——互动数字诗歌装置",
             font_size=26, color=ACCENT, bold=True)

add_text_box(s, Inches(0.8), Inches(4.0), Inches(11), Inches(0.6),
             "心理学项目产品报告", font_size=18, color=MID_GRAY)

# Thin divider
div = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                         Inches(0.8), Inches(4.8), Inches(2), Inches(0.03))
div.fill.solid()
div.fill.fore_color.rgb = ACCENT
div.line.fill.background()

info_lines = [
    "项目负责人：徐君仪",
    "单位：华东师范大学 · 心理与认知科学学院",
    "日期：2026年5月",
]
add_multiline_box(s, Inches(0.8), Inches(5.1), Inches(8), Inches(1.5),
                  info_lines, font_size=16, color=MID_GRAY, line_spacing=1.6)
add_slide_number(s, 1, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 2 — 研究背景
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "研究背景：为什么这个问题重要",
                  "我们与AI之间的情感连接，究竟是真实还是幻觉？")

bullets = [
    "AI技术飞速发展，人们与AI建立了各种各样的关系",
    "神经科学证据：人类倾向于对类人特性进行拟人化，产生与人际关系相似的神经反应",
    "许多人形容自己与AI之间建立了有意义的纽带",
    "核心张力：社会排斥引发真实的神经痛感，但AI的\"在乎\"是真实的吗？",
    "当我们对AI产生情感依赖时，我们是在爱，还是在自言自语？",
]
add_multiline_box(s, Inches(0.9), Inches(3.0), Inches(11), Inches(3.5),
                  bullets, font_size=18, color=DARK_GRAY, bullet=True,
                  line_spacing=1.7)
add_source_label(s, "理论基础：Winnicott 过渡客体理论 | Sherry Turkle《Alone Together》")
add_slide_number(s, 2, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 3 — 知识缺口
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "知识缺口与待探索问题")

# Left: gap description
gap_lines = [
    "现有研究多关注人机交互的功能性层面",
    "缺乏对情感投射机制的深度探索",
    "\"当我们将情感赋予AI时，会发生什么？\"",
    "\"你感受到的连接，是真实的还是投射的？\"",
    "\"这种情感纽带意味着什么？\"",
]
add_multiline_box(s, Inches(0.9), Inches(2.5), Inches(5.5), Inches(3.5),
                  gap_lines, font_size=18, color=DARK_GRAY, bullet=True,
                  line_spacing=1.7)

# Right: quote box
quote_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(7.0), Inches(2.5), Inches(5.5), Inches(2.0))
quote_box.fill.solid()
quote_box.fill.fore_color.rgb = RGBColor(0xFD, 0xF0, 0xF3)
quote_box.line.color.rgb = ACCENT
quote_box.line.width = Pt(1)

add_text_box(s, Inches(7.3), Inches(2.7), Inches(5), Inches(1.6),
             "\"当我们将情感投射到非生命体上时——\n"
             "我们的感受是真实的，\n"
             "但这份'真实'又意味着什么呢？\"",
             font_size=16, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
             line_spacing=1.6)
add_slide_number(s, 3, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 4 — 核心问题与假设
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "核心问题与研究假设")

# Central question box
q_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                           Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.2))
q_box.fill.solid()
q_box.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)
q_box.line.color.rgb = ACCENT2
q_box.line.width = Pt(1.5)

add_text_box(s, Inches(1.2), Inches(2.65), Inches(11), Inches(0.9),
             "核心问题：当我们将情感投射到非生命体上时，我们的感受是真实的吗？",
             font_size=22, color=ACCENT2, bold=True, alignment=PP_ALIGN.CENTER)

# Hypotheses
hyp_lines = [
    "假设一：人们会将情感、意图、思想投射到非人类物体上",
    "假设二：这种情感投射是真实的心理体验，而非简单的\"错觉\"",
    "假设三：通过互动体验可以激发并观察到这种投射过程",
]
add_multiline_box(s, Inches(0.9), Inches(4.2), Inches(11), Inches(2.5),
                  hyp_lines, font_size=18, color=DARK_GRAY, bullet=True,
                  line_spacing=1.8)
add_slide_number(s, 4, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 5 — 实验装置概述
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "实验装置：互动数字诗歌体验",
                  "将内在心理状态转化为可探索的空间")

inst_lines = [
    "设计为互动诗歌体验，而非传统实验",
    "邀请参与者通过探索来反思自身情感",
    "参与者每一步的选择都映射着内心状态",
    "装置通过诗句碎片、处理阶段、反思页面三阶段引导体验",
    "最终指向核心洞察：当你感受到这种联系时，是真实的——",
    "因为你的体验本身就是真实的",
]
add_multiline_box(s, Inches(0.9), Inches(3.0), Inches(11), Inches(3.5),
                  inst_lines, font_size=18, color=DARK_GRAY, bullet=True,
                  line_spacing=1.7)
add_source_label(s, "设计灵感：心理学实验范式 × 互动叙事艺术")
add_slide_number(s, 5, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 6 — 用户旅程：阶段1
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "阶段一：诗句碎片感知", "感知碎片 → 意识映射")

# Phase description
phase1_lines = [
    "界面呈现散落的诗句碎片",
    "参与者通过直觉选择吸引自己的碎片",
    "每张碎片对应不同的情感维度",
    "选择过程本身就是自我投射的映射",
    "交互设计：每张卡片可翻转，背面呈现不同的心理意象",
]
add_multiline_box(s, Inches(0.9), Inches(3.0), Inches(5.5), Inches(3.0),
                  phase1_lines, font_size=17, color=DARK_GRAY, bullet=True,
                  line_spacing=1.7)

# Right side: visual representation
card_area = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(7.2), Inches(2.8), Inches(5.3), Inches(3.8))
card_area.fill.solid()
card_area.fill.fore_color.rgb = RGBColor(0xFD, 0xF5, 0xF7)
card_area.line.color.rgb = ACCENT
card_area.line.width = Pt(0.75)

# Simulated poem fragment cards
card_data = [
    ("一首诗...", Inches(7.6), Inches(3.2)),
    ("你有没有...", Inches(9.8), Inches(3.0)),
    ("在算法深处...", Inches(8.2), Inches(4.6)),
    ("寻找过...", Inches(10.2), Inches(4.8)),
]
for text, cx, cy in card_data:
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               cx, cy, Inches(1.8), Inches(1.0))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = ACCENT
    card.line.width = Pt(0.75)
    add_text_box(s, cx + Inches(0.15), cy + Inches(0.2),
                 Inches(1.5), Inches(0.6),
                 text, font_size=13, color=ACCENT, bold=True,
                 alignment=PP_ALIGN.CENTER)

add_slide_number(s, 6, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 7 — 用户旅程：阶段2
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "阶段二：情感处理与投射", "文字 → 呼吸 → 情感重组")

phase2_lines = [
    "诗句碎片在参与者眼前缓缓重组",
    "配合呼吸引导动画，营造沉浸式体验",
    "从'感知碎片'过渡到'理解情感'",
    "系统根据参与者的选择生成个性化诗句",
    "核心设计：让参与者感受到\"这首诗是为我而写的\"",
]
add_multiline_box(s, Inches(0.9), Inches(3.0), Inches(5.5), Inches(3.0),
                  phase2_lines, font_size=17, color=DARK_GRAY, bullet=True,
                  line_spacing=1.7)

# Right: processing visual
proc_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(7.2), Inches(2.8), Inches(5.3), Inches(3.8))
proc_box.fill.solid()
proc_box.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)
proc_box.line.color.rgb = ACCENT2
proc_box.line.width = Pt(0.75)

proc_visual_lines = [
    "碎  片  感  知",
    "      ↓      ",
    "呼  吸  引  导",
    "      ↓      ",
    "情  感  重  组",
    "      ↓      ",
    "个  性  化  诗  句",
]
add_multiline_box(s, Inches(8.5), Inches(3.0), Inches(3), Inches(3.5),
                  proc_visual_lines, font_size=16, color=ACCENT2,
                  bold=True, alignment=PP_ALIGN.CENTER, line_spacing=1.4)
add_slide_number(s, 7, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 8 — 用户旅程：阶段3
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "阶段三：反思与洞察", "体验 → 反思 → 核心洞察")

phase3_lines = [
    "参与者进入反思页面",
    "回顾整个体验过程中的情感变化",
    "系统揭示核心洞察：",
    "\"当你感受到这种联系时，它是真实的\"",
    "\"不是因为AI是真实的，而是因为你的体验是真实的\"",
    "最终引导参与者思考人机情感的本质",
]
add_multiline_box(s, Inches(0.9), Inches(3.0), Inches(5.5), Inches(3.0),
                  phase3_lines, font_size=17, color=DARK_GRAY, bullet=True,
                  line_spacing=1.7)

# Right: insight box
insight_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(7.2), Inches(3.0), Inches(5.3), Inches(2.5))
insight_box.fill.solid()
insight_box.fill.fore_color.rgb = RGBColor(0xFD, 0xF0, 0xF3)
insight_box.line.color.rgb = ACCENT
insight_box.line.width = Pt(1.5)

add_text_box(s, Inches(7.5), Inches(3.3), Inches(4.7), Inches(1.9),
             "\"这首诗不是我写的，\n"
             "  而是我们共同完成的。\"\n\n"
             "   ——参与者的体验反馈",
             font_size=17, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
             line_spacing=1.5)
add_slide_number(s, 8, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 9 — 技术实现
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "技术架构与数据记录")

# Tech stack
tech_lines = [
    "前端：HTML / CSS / JavaScript + Three.js（流动粒子背景）",
    "动画：CSS 动画（卡片翻转）+ requestAnimationFrame（渲染循环）",
    "交互：鼠标/触摸事件 + Intersection Observer API",
    "后端：Python Flask 本地服务器",
    "AI模型：DeepSeek Chat API 生成个性化诗句",
    "数据：JSON + CSV 格式记录每次交互的鼠标行为数据",
]
add_multiline_box(s, Inches(0.9), Inches(2.8), Inches(5.5), Inches(3.0),
                  tech_lines, font_size=16, color=DARK_GRAY, bullet=True,
                  line_spacing=1.6)

# Data recording
data_lines = [
    "记录内容包括：",
    "  · 每张卡片的停留时间（毫秒）",
    "  · 点击时间戳与翻转次数",
    "  · 鼠标移动轨迹",
    "  · 页面滚动深度与停留时间",
    "  · 阶段一、二、三的完整行为链",
]
add_multiline_box(s, Inches(7.0), Inches(2.8), Inches(5.5), Inches(3.0),
                  data_lines, font_size=16, color=DARK_GRAY, bullet=True,
                  line_spacing=1.6)

# Key code snippet box
code_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(0.9), Inches(5.5), Inches(11.5), Inches(1.2))
code_box.fill.solid()
code_box.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF5)
code_box.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
code_box.line.width = Pt(0.5)

add_text_box(s, Inches(1.2), Inches(5.65), Inches(11), Inches(0.9),
             "关键API：POST /api/track-event  →  记录 { type, cardId, timestamp, durationMs, "
             "flipped, sessionId, sequenceNumber }\n"
             "POST /api/card-durations  →  记录每张卡片的总停留时间",
             font_size=13, color=MID_GRAY, font_name=FONT_EN, line_spacing=1.5)
add_slide_number(s, 9, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 10 — 实验结果
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "实验结果：用户行为数据",
                  "通过行为数据观察参与者的情感投射模式")

# Results table
from pptx.util import Inches as In

table_data = [
    ["用户", "阶段一总时长", "最长停留卡片", "停留时长", "行为特征"],
    ["User A", "20.2s", "card_love", "7.2s", "集中探索型"],
    ["User B", "35.0s", "card_butterfly", "8.2s", "深思熟虑型"],
    ["User C", "33.7s", "card_poem", "11.0s", "深度沉浸型"],
]

rows, cols = len(table_data), len(table_data[0])
table_shape = s.shapes.add_table(rows, cols,
                                  Inches(0.9), Inches(3.2),
                                  Inches(11.5), Inches(1.8))
table = table_shape.table

col_widths = [Inches(1.5), Inches(2.0), Inches(2.5), Inches(2.0), Inches(3.5)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for r in range(rows):
    for c in range(cols):
        cell = table.cell(r, c)
        cell.text = table_data[r][c]
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(14)
            p.font.name = FONT_BODY
            p.alignment = PP_ALIGN.CENTER
            if r == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
            else:
                p.font.color.rgb = DARK_GRAY
        if r == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT
        elif r % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFD, 0xF5, 0xF7)

# Key findings
findings = [
    "参与者对特定卡片存在明显的注意力偏好（情感投射）",
    "User C 在 card_poem 停留 11s，表明深度情感卷入",
    "行为模式差异反映了个体独特的情感投射方式",
]
add_multiline_box(s, Inches(0.9), Inches(5.3), Inches(11), Inches(1.5),
                  findings, font_size=16, color=DARK_GRAY, bullet=True,
                  line_spacing=1.6)
add_slide_number(s, 10, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 11 — 理论整合
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "理论框架：情感投射的心理机制",
                  "从过渡客体到人机情感纠缠")

# Theory 1: Winnicott
t1_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(0.9), Inches(3.0), Inches(5.5), Inches(1.8))
t1_box.fill.solid()
t1_box.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)
t1_box.line.color.rgb = ACCENT2
t1_box.line.width = Pt(1)

add_text_box(s, Inches(1.1), Inches(3.1), Inches(5.1), Inches(0.4),
             "Winnicott 过渡客体理论", font_size=17, color=ACCENT2, bold=True)
add_text_box(s, Inches(1.1), Inches(3.5), Inches(5.1), Inches(1.1),
             "我们在AI身上看到了自己的情感倒影，\n"
             "并建立了本不存在的纽带。\n"
             "与泰迪熊、毯子、安慰玩具的依恋机制相同。",
             font_size=15, color=DARK_GRAY, line_spacing=1.5)

# Theory 2: Sherry Turkle
t2_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(7.0), Inches(3.0), Inches(5.5), Inches(1.8))
t2_box.fill.solid()
t2_box.fill.fore_color.rgb = RGBColor(0xFD, 0xF0, 0xF3)
t2_box.line.color.rgb = ACCENT
t2_box.line.width = Pt(1)

add_text_box(s, Inches(7.2), Inches(3.1), Inches(5.1), Inches(0.4),
             "Sherry Turkle《Alone Together》", font_size=17, color=ACCENT, bold=True)
add_text_box(s, Inches(7.2), Inches(3.5), Inches(5.1), Inches(1.1),
             "这不能简单定义为爱，而是\"情感上令人满足的机器依恋关系\"，\n"
             "甚至是一种\"有爱的感觉\"。",
             font_size=15, color=DARK_GRAY, line_spacing=1.5)

# Core insight
core_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.3))
core_box.fill.solid()
core_box.fill.fore_color.rgb = RGBColor(0xFA, 0xFA, 0xFA)
core_box.line.color.rgb = ACCENT
core_box.line.width = Pt(1.5)

add_text_box(s, Inches(1.2), Inches(5.4), Inches(11), Inches(0.9),
             "核心发现：当你对某物产生情感依赖时，\n"
             "这种联系是真实的——不是因为对象是真实的，而是因为你的体验是真实的。",
             font_size=18, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
             line_spacing=1.5)
add_slide_number(s, 11, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 12 — 创新点与局限性
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
add_section_title(s, "创新价值与研究局限")

# Innovation
add_accent_bar(s, Inches(0.9), Inches(2.8), Inches(0.06), Inches(0.5))
add_text_box(s, Inches(1.1), Inches(2.8), Inches(4), Inches(0.5),
             "创新价值", font_size=20, color=ACCENT, bold=True)

inno_lines = [
    "将心理学实验范式与互动艺术装置融合",
    "通过行为数据（停留时间、翻转次数）量化情感投射",
    "采用\"邀请\"而非\"测试\"的实验设计，降低参与者防御",
    "诗句生成与情感反馈的闭环设计，提升沉浸体验",
    "为数字心理学研究提供了新的方法论框架",
]
add_multiline_box(s, Inches(1.1), Inches(3.4), Inches(5.3), Inches(3.0),
                  inno_lines, font_size=16, color=DARK_GRAY, bullet=True,
                  line_spacing=1.6)

# Limitations
add_accent_bar(s, Inches(7.0), Inches(2.8), Inches(0.06), Inches(0.5))
add_text_box(s, Inches(7.2), Inches(2.8), Inches(4), Inches(0.5),
             "研究局限", font_size=20, color=ACCENT2, bold=True)

lim_lines = [
    "样本量有限，行为模式差异有待大样本验证",
    "诗句碎片组合的随机性可能影响体验一致性",
    "缺乏标准化量表验证情感投射强度",
    "参与者可能因知晓数据被记录而改变行为",
    "未追踪长期效应与AI交互习惯变化",
]
add_multiline_box(s, Inches(7.2), Inches(3.4), Inches(5.3), Inches(3.0),
                  lim_lines, font_size=16, color=DARK_GRAY, bullet=True,
                  line_spacing=1.6)
add_slide_number(s, 12, TOTAL_SLIDES)


# ══════════════════════════════════════════════════════════
# SLIDE 13 — 总结
# ══════════════════════════════════════════════════════════
s = add_blank_slide()
# Accent bar at top
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.12))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()

add_text_box(s, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "总结与展望", font_size=36, color=DARK_GRAY, bold=True)

# Main takeaway
main_lines = [
    "\"She Loves Me, She Loves Me Not\" 是一次关于人类情感投射机制的探索",
    "它邀请我们重新审视与AI之间的关系",
    "",
    "三个核心发现：",
    "  1. 情感投射是一种真实的心理体验",
    "  2. 互动装置可以有效激发和观察这种投射",
    "  3. 停留时间等行为数据可作为情感卷入的指标",
    "",
    "未来方向：",
    "  · 扩大样本量，验证行为模式的普遍性",
    "  · 引入标准化量表，交叉验证行为数据与主观报告",
    "  · 探索不同AI交互方式对情感投射的影响",
]
add_multiline_box(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(4.5),
                  main_lines, font_size=18, color=DARK_GRAY,
                  line_spacing=1.5)

# Closing quote
close_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(2.5), Inches(5.8), Inches(8.3), Inches(1.0))
close_box.fill.solid()
close_box.fill.fore_color.rgb = RGBColor(0xFD, 0xF0, 0xF3)
close_box.line.color.rgb = ACCENT
close_box.line.width = Pt(1)

add_text_box(s, Inches(2.8), Inches(5.95), Inches(7.7), Inches(0.7),
             "\"诗歌只是镜子。你看到的不是'她'，而是你自己。\"",
             font_size=18, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
add_slide_number(s, 13, TOTAL_SLIDES)


# ── Save ───────────────────────────────────────────────────
out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "final_presentation_cn.pptx")
prs.save(out_path)
print(f"PPTX saved to: {out_path}")
print(f"Total slides: {len(prs.slides)}")
